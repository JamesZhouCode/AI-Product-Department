# -*- coding: utf-8 -*-
"""
从「环节与上市企业入链v2.xlsx」+「储能产业链_头部企业标记.xlsx」抽取储能产业数据，
生成两份前端产物：

  1) public/data/industries/storage-chain.json   —— 流程页（建链四步）用
  2) public/data/industries/companies-storage.json —— 「储能产业」一级链数据分片
     （供 图谱 / 地图 / 周报 页签复用，需满足 test/newIndustryData.test.js 的契约）

关键约定：
- **Sheet1 是权威节点框架**（位置 / 一级环节 / 二级环节 / 顺序），顺序 1..24 决定并排次序。
- 节点用 **(一级环节, 二级环节) 复合键**：二级环节名「检测认证与标准」出现两次
  （顺序 8 属一级「检测认证与标准」；顺序 24 属一级「应用与运营」），单用名字会互相覆盖。
- **顺序只用于排序，不渲染到界面**（用户明确要求二级环节前不显示序号）。
- 头部标签（龙头/骨干）来自头部标记表的 `01_分环节头部企业` sheet，
  其「环节」列比二级环节更细（如「正极材料/负极材料/电解液」都属「电池主材与辅材」），
  因此保留为 `heads[].segment` 细分环节分组。
- 未出现在头部名单中的企业记为**长尾**；名单覆盖全部 237 家（头部 176 + 长尾 61）。
- 分档规则整体落在 `headRule` 配置对象里，阈值/权重/门槛均可改配置，无需改代码。

重新运行：
  python frontend/scripts/extract-storage-chain.py
"""
import json
import os
import re
import hashlib
import collections

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_CHAIN = r"E:\产业链\电力装备\储能所需数据\环节与上市企业入链v2.xlsx"
SRC_HEADS = r"E:\产业链\电力装备\储能所需数据\储能产业链_头部企业标记.xlsx"
SRC_REL = r"E:\产业链\电力装备\储能所需数据\上市企业关系档案_237家_含属地.xlsx"
# 全量客户清单：节点关系（含官方二级环节编码）+ 237 上市客户 + 非 237 的其他客户
SRC_ALL = r"E:\产业链\电力装备\储能所需数据\全量客户清单（含上市公司）.xlsx"
OUT_CHAIN = os.path.join(BASE, "public", "data", "industries", "storage-chain.json")
OUT_SHARD = os.path.join(BASE, "public", "data", "industries", "companies-storage.json")

# 落点数据源（全部本地文件，运行时零外部请求）
# ① 真实坐标缓存：仓库里既有的一次性高德研究产物，按企业全称命中即用
LOCATION_CACHE_FILES = [
    "docs/new-industry-geocode-research.json",
    "docs/coordinate-research.json",
    "docs/listed-location-candidates.json",
]
# ② 地名词表：从本地 PMTiles 底图抽出的 OSM 地名点，用于「区域参考」落点
GAZETTEER_FILE = "docs/place-gazetteer.json"
# ③ 省界：给每一个坐标做「是否落在标称省份内」的校验。
#    词表里存在大量同名地名（跨省的「朝阳区」「和平区」，甚至跨国的同名点），
#    不校验省份就会把企业放到错误省份 —— 实际发生过：广州 60 家企业被放到韩国。
PROVINCE_GEOJSON = os.path.join(BASE, "public", "data", "china-provinces.geojson")

INDUSTRY_LABEL = "储能产业"
INDUSTRY_CODE = "ES"

# 一级环节的展示说明
LEVEL1_NOTE = {
    "资源与材料": "矿产 · 主材辅材 · 回收",
    "核心电子器件": "电力电子与功率器件",
    "核心装备": "智造装备 · 长时储能 · 氢储",
    "检测认证与标准": "并网检测 · 标准认证",
    "电池与核心部件": "电芯 · PACK · BMS · EMS · PCS · 热管理",
    "系统集成与工程建设": "集成 · EPC · 链主 · 三方集成",
    "应用与运营": "电源/电网/用户侧 · 运营服务",
}

LEVEL2_NOTE = {
    "矿产资源与化工原料": "锂/钴/镍/磷/钒等上游资源",
    "电池主材与辅材": "正负极、隔膜、电解液、结构件",
    "回收与循环利用": "退役电池梯次与再生",
    "关键电力电子与器件": "IGBT、MOSFET、电容器等",
    "物理/长时储能装备": "压缩空气、飞轮、液流、熔盐",
    "氢储能装备": "电解槽与储氢",
    "智能制造装备": "涂布、卷绕、化成分容设备",
    "检测认证与标准": "型式试验、并网检测、认证",
    "储能电池本体": "大储/工商业/户储电芯与电池",
    "电池PACK与模组": "模组与 PACK 集成",
    "电池管理系统BMS": "电芯均衡、状态估计",
    "能量管理系统 EMS": "调度策略、能量优化",
    "储能变流器 PCS": "双向变流、构网型",
    "热管理与安全防护": "液冷、消防、结构安全",
    "储能系统集成": "直流侧+交流侧一体化集成",
    "EPC与建安": "工程总承包与安装调试",
    "集采链主/生态主导型": "央国企集采与生态主导方",
    "第三方系统集成商": "独立第三方系统交付",
    "电源侧应用": "风光配储、独立储能电站",
    "电网侧应用": "调峰调频、共享储能",
    "用户侧应用": "工商业与户用储能",
    "新场景与跨业态": "数据中心、充电站等",
    "运营服务与电力市场": "运维、容量租赁、电力交易",
}

POSITION_NOTE = {
    "上游": "资源 · 器件 · 装备 · 检测",
    "中游": "电池与核心部件 · 系统集成",
    "下游": "应用场景 · 运营服务",
}

GRADE_ORDER = ["核心主业", "核心", "相关重要", "边缘观察", "关联"]

# ---------------- 评分明细（占位实现） ----------------
# ⚠️ 上游底表目前**没有**逐企业的评分明细，这里按已知档位反推一份**占位评分**，
#    目的是让 step3 的「配置维度/权重 → 重新评估分档」这条链路先跑通、可演示。
#    接入真实评分表时，只要把这里替换成读取该表的四个维度分即可，前端无需改动
#    （前端只认 scoreDetail 的四个键与权重配置）。
#
# 维度分两组：
#   base（基础维度）= 量：份额 / 规模 / 成长 —— 回答「现在有多大」
#   moat（壁垒维度）= 质：技术 / 资质 / 客户 / 卡位 / 成本 —— 回答「领先守不守得住」
# 默认只启用 base 3 项 + tech（合计 100），与既有口径一致；其余壁垒维度默认权重 0、
# 处于「未启用」状态，供界面上一键新增后参与评估。
#
# 生成方式：先按档位取一个目标总分 T，再在默认启用的四个维度上做确定性扰动（以企业名
# 做种子，保证每次生成完全一致），并让「加权总分」在扰动前后保持不变 —— 所以默认配置下
# 与上游名单 100% 一致。壁垒维度以 tech 分为锚做扰动，使同类企业的壁垒水平彼此自洽。
SCORE_DIMENSIONS = ["share", "tech", "scale", "growth"]  # 默认启用的四项
SCORE_WEIGHTS_DEFAULT = {"share": 40, "tech": 25, "scale": 20, "growth": 15}
SCORE_ANCHOR = {
    "leader": (83, 92),  # 龙头：默认权重下稳定落在 ≥80
    "core": (63, 77),  # 骨干：60–79
    "tail": (42, 58),  # 长尾：<60
}
SCORE_DRIFT = 16  # 单维度偏离目标总分的最大幅度
MOAT_DRIFT = 10  # 壁垒维度相对 tech 锚点的最大扰动
SCORE_MIN, SCORE_MAX = 5, 99

# 维度池：kind = base / moat，weight = 默认权重，enabled = 是否默认参与评估。
# 顺序即界面展示顺序（先基础维度，再壁垒维度）。
SCORE_DIMENSION_POOL = [
    {
        "key": "share",
        "label": "国内份额与排名",
        "kind": "base",
        "weight": 40,
        "enabled": True,
        "note": "区分大储 / 工商储、动力与储能口径",
    },
    {
        "key": "scale",
        "label": "规模与盈利",
        "kind": "base",
        "weight": 20,
        "enabled": True,
        "note": "营收规模、产能与毛利率",
    },
    {
        "key": "growth",
        "label": "客户结构与成长",
        "kind": "base",
        "weight": 15,
        "enabled": True,
        "note": "头部客户占比与订单成长性",
    },
    {
        "key": "tech",
        "label": "技术壁垒",
        "kind": "moat",
        "weight": 25,
        "enabled": True,
        "note": "核心工艺、专利与标准话语权",
    },
    {
        "key": "cert",
        "label": "资质与认证壁垒",
        "kind": "moat",
        "weight": 0,
        "enabled": False,
        "note": "并网 / 安规 / 消防等强制准入，央国企集采入库",
    },
    {
        "key": "customer",
        "label": "客户壁垒",
        "kind": "moat",
        "weight": 0,
        "enabled": False,
        "note": "头部客户切换成本、长协与联合开发绑定",
    },
    {
        "key": "position",
        "label": "产业链卡位壁垒",
        "kind": "moat",
        "weight": 0,
        "enabled": False,
        "note": "关键环节稀缺卡位、上下游一体化程度",
    },
    {
        "key": "cost",
        "label": "一体化与成本壁垒",
        "kind": "moat",
        "weight": 0,
        "enabled": False,
        "note": "自供率、规模成本曲线与单位成本优势",
    },
]

DIMENSION_KIND = {d["key"]: d["kind"] for d in SCORE_DIMENSION_POOL}
MOAT_DIMENSIONS = [d["key"] for d in SCORE_DIMENSION_POOL if d["kind"] == "moat"]


def _seed01(name, salt):
    """用企业名 + 维度名生成 0..1 的确定性伪随机数（不用 random，保证可复现）。"""
    digest = hashlib.md5(f"{name}|{salt}".encode("utf-8")).digest()
    return int.from_bytes(digest[:4], "big") / 2**32


def build_score_detail(name, tier):
    """按档位反推一份占位评分明细。默认启用维度的加权总分严格等于目标总分（钳位前）。

    返回全部 8 个维度的分（含默认未启用的壁垒维度），这样界面上「新增维度」时
    企业侧有分可用，不必回头改生成逻辑。
    """
    lo, hi = SCORE_ANCHOR.get(tier, SCORE_ANCHOR["tail"])
    target = lo + (hi - lo) * _seed01(name, "total")

    raw = {d: (_seed01(name, d) - 0.5) * 2 * SCORE_DRIFT for d in SCORE_DIMENSIONS}
    # 扣掉加权均值，使 Σ(w·score)/Σw 仍等于 target
    weight_sum = sum(SCORE_WEIGHTS_DEFAULT[d] for d in SCORE_DIMENSIONS)
    weighted_mean = (
        sum(SCORE_WEIGHTS_DEFAULT[d] * raw[d] for d in SCORE_DIMENSIONS) / weight_sum
    )
    detail = {}
    for d in SCORE_DIMENSIONS:
        value = target + (raw[d] - weighted_mean)
        detail[d] = int(round(min(SCORE_MAX, max(SCORE_MIN, value))))

    # 壁垒维度以 tech 为锚点做小扰动：一家企业的技术壁垒高，其余壁垒通常也不弱，
    # 避免出现「技术 90 分、客户壁垒 30 分」这类自相矛盾的组合。
    anchor = detail["tech"]
    for d in MOAT_DIMENSIONS:
        if d in detail:
            continue
        value = anchor + (_seed01(name, d) - 0.5) * 2 * MOAT_DRIFT
        detail[d] = int(round(min(SCORE_MAX, max(SCORE_MIN, value))))
    return detail

# 企业简称在两份表里的写法差异（均为更名/别称，非编造）
NAME_ALIASES = {
    "五矿新能": "长远锂科",  # 688779 更名
    "青鸟智控": "青鸟消防",  # 002960 更名
}

# 头部名单里形如「赣锋锂业(龙头)、天齐锂业(龙头)」
HEAD_PATTERN = re.compile(r"([^、（）()]+?)[（(](龙头|骨干)[）)]")

TIER_LABEL = {"leader": "龙头", "core": "骨干", "tail": "长尾"}

# ---------------- 链外被点名主体：抽取规则 ----------------
# 只看「前五大客户 / 前五大供应商」两栏（「主要上下游配套」多为行业描述，噪音太大）。
EXTERNAL_FIELDS = [("前五大供应商", "upstream"), ("前五大客户", "downstream")]
EXTERNAL_MIN_COUNT = 2  # 只留被 >=2 次点名的主体，压住长尾噪音
EXTERNAL_MAX_SOURCES = 8  # 每个主体最多记 8 个点名来源

# 名称尾部常见的虚词（「等」「等（锂盐、前驱体）」之类）先剥掉
EXTERNAL_TRAIL = re.compile(r"[等及其他与和]+$")
# 口径词 / 行业词 / 文本片段：这些不是主体名，直接判掉
EXTERNAL_NOISE = re.compile(
    r"^(未披露|未公开|未具名|未指明|年报未|天眼查|工商记录|公开记录|公开披露|公开资料|公开报道"
    r"|以['\"]?[客供]户|具体供应商名称|其中|另有|另含|以及|包括|主要为|下游|上游)"
)
EXTERNAL_SUFFIX_NOISE = re.compile(
    r"(原料供应商|材料供应商|设备供应商|及设备供应商|供应商|厂商|企业|客户群)$"
)
# 纯物料/行业类词（非主体）
EXTERNAL_MATERIAL_NOISE = {
    "锂盐",
    "碳酸锂",
    "氢氧化锂",
    "磷酸铁",
    "石油焦",
    "溶剂",
    "钢材",
    "铜材",
    "铝材",
    "芯片",
    "硅片",
    "结构件",
    "铸锻件",
    "压缩机",
    "光伏组件",
    "正极材料",
    "负极材料",
    "电解液",
    "隔膜",
}
# 归一：同一主体的不同写法（只处理确凿的同义写法）
EXTERNAL_NORMALIZE = {
    "国家电网有限公司": "国家电网",
    "国家电网及地方电网公司": "国家电网",
    "中国南方电网有限责任公司": "南方电网",
    "南方电网各省市电力公司": "南方电网",
    "国家电投": "国家电投集团",
    "阿里": "阿里巴巴",
    "Lenovo": "联想",
    "长安": "长安汽车",
    "华电": "华电集团",
    "金川": "金川集团",
}
# 明确剔除（描述性统称或文本碎片，不宜当作主体）
EXTERNAL_DROP = {"国家电网及地方电网公司"}

# 人工核实补充：抽取门槛（>=2 次点名）会漏掉只被点名一次、但确属本链环节的非上市主体。
# 这里单独记一条，避免「系统没抽到」被误读成「链里没有」。事实均有公开来源，不做推测。
EXTERNAL_VERIFIED = [
    {
        "name": "河北吉诚新材料有限公司",
        "short": "河北吉诚",
        "side": "上游",
        "segment": "资源与材料",
        "listed": False,
        "note": (
            "非上市（其他有限责任公司）。主营电池级碳酸锂 / 氢氧化锂，一期各 1 万吨 / 年产能，"
            "另有氯化法钛白粉、针状焦等；河北燕山钢铁集团有限公司持股 70%。"
        ),
        "evidence": "天齐锂业在「前五大供应商」中具名披露；工商登记与公开债券披露材料可查。",
        "mentionedBy": ["天齐锂业"],
        "source": "公开工商登记 / 企业公开披露",
    },
]


def looks_like_external_entity(tok):
    """判断一个切片是否是「像主体的名字」，用来滤掉口径词 / 行业词 / 文本碎片。"""
    if len(tok) < 2 or tok in EXTERNAL_MATERIAL_NOISE:
        return False
    if EXTERNAL_NOISE.match(tok) or EXTERNAL_SUFFIX_NOISE.search(tok):
        return False
    if re.match(r"^(19|20)\d{2}年?", tok):
        return False
    if re.fullmatch(r"[\d\.,%万亿元亿千瓦时吨]+", tok):
        return False
    if not re.search(r"[\u4e00-\u9fa5A-Za-z]", tok):
        return False
    # 至少要有 2 个连续汉字，或 3 个以上拉丁字母（品牌名）
    if not re.search(r"[\u4e00-\u9fa5]{2,}|[A-Za-z]{3,}", tok):
        return False
    if "%" in tok or "亿元" in tok or "万元" in tok or "列示" in tok or "代称" in tok:
        return False
    # 含未配对的引号 → 多为引文碎片
    if tok.count('"') or tok.count("”") or tok.count("“") or tok.count("'"):
        return False
    return True


# ---------------- 分档规则（可配置） ----------------
HEAD_RULE = {
    "version": "2026-09-10",
    "label": "储能产业头部企业分档规则",
    "coverage": "237 家上市企业（A 股 + 港股），按环节归属逐一评估",
    "basis": "以国内份额为主（区分大储/工商储、动力与储能口径）",
    "scoreTotal": 100,
    # 维度分两组：基础维度（量）与壁垒维度（质 / 护城河）。
    # 界面按此渲染，并支持新增 / 删减维度后实时重算 —— 改这里即可调整维度池。
    "dimensionGroups": [
        {"key": "base", "label": "基础维度", "note": "量：企业现在有多大"},
        {"key": "moat", "label": "壁垒维度", "note": "质：领先能不能守住"},
    ],
    "dimensions": SCORE_DIMENSION_POOL,
    # 壁垒门槛（默认关闭）：总分够龙头但壁垒分不足时降为骨干。
    # 这是「建立壁垒」的落点 —— 龙头不仅要大，还要守得住。
    "moatGate": {
        "key": "moatGate",
        "enabled": False,
        "min": 70,
        "label": "壁垒分门槛",
        "desc": "启用后：总分达到龙头阈值、但壁垒分低于门槛的企业自动降为骨干，避免「大而不强」被误判为龙头。",
    },
    "tiers": [
        {
            "key": "leader",
            "label": "龙头",
            "min": 80,
            "max": None,
            "desc": "综合得分 ≥80，且通过全部硬门槛",
        },
        {"key": "core", "label": "骨干", "min": 60, "max": 79, "desc": "综合得分 60–79"},
        {"key": "tail", "label": "长尾", "min": 0, "max": 59, "desc": "综合得分 <60"},
    ],
    "hardGates": [
        "主业确属该环节",
        "有可查的国内份额 / 排名或头部客户证据",
        "具备该环节核心资质 / 认证之一",
        "至少具备一项可守的壁垒（技术 / 资质 / 客户 / 卡位 / 成本）",
    ],
    "sources": [
        "GGII 高工",
        "鑫椤锂电",
        "CNESA / 中关村储能联盟",
        "中国光伏行业协会（CPIA）",
        "风能协会（CWEA）",
        "上市公司年报 / 公告",
        "券商研报",
        "招投标与行业新闻",
    ],
    "limitations": (
        "份额类数据多为付费或半年更新；无公开份额时以可查证据（出货 / 中标 / 客户 / 产能）"
        "做保守判断，依据列标注来源与年份。"
    ),
    "knownIssues": [
        {"topic": "长时储能", "note": "名单内无该环节明确龙头（行业龙头多为非上市，如大连融科、星辰新能），最高为骨干。"},
        {"topic": "消防 / 安全", "note": "青鸟智控按「储能消防」归入温控 / 热管理，如需可单设「安全消防」节点。"},
        {
            "topic": "光伏 / 风电设备",
            "note": "含特变电工、东方电气、中天科技等电力装备全口径企业，储能关联见依据列。",
        },
        {"topic": "环节归属", "note": "关键词自动分级 + 各环节簇人工校正，个别边界案例可按业务口径再调。"},
    ],
    # 说明哪些字段可改配置；改这里即可调整分档，不需要改前端代码
    "configurable": {
        "path": "headRule",
        "editable": [
            "dimensions[].weight",
            "dimensions[].enabled（新增 / 删减维度）",
            "tiers[].min / tiers[].max",
            "moatGate.enabled / moatGate.min",
            "hardGates",
            "basis",
            "sources",
        ],
        "note": (
            "维度池、权重、阈值、壁垒门槛都在此处配置；前端只负责渲染与实时重算，"
            "调整本对象即可改变评估口径，不需要改前端代码。"
        ),
    },
    # 评分明细的来源说明。上游底表暂无逐企业评分，scoreDetail 为按已知档位反推的占位数据；
    # 接入真实评分表后把 simulated 置 false 并替换生成逻辑即可，前端无需改动。
    "scoreSource": {
        "simulated": True,
        "note": (
            "上游底表当前没有逐企业评分明细，这里各维度的分为按已知档位反推的占位值，"
            "用于演示「新增 / 删减维度、调权重、设壁垒门槛 → 重新评估分档」的完整链路；"
            "接入真实评分表后替换数据即可，评估逻辑不变。"
        ),
        "anchorNote": "占位评分按档位锚定：龙头 83–92、骨干 63–77、长尾 42–58（默认权重下总分即落在此区间）。",
        "moatNote": "壁垒维度分以技术壁垒为锚点生成，同一家企业的各项壁垒分互相自洽。",
    },
}


def norm(value):
    return str(value).strip() if value is not None else ""


# ---------------- 落点：本地缓存 + 地名词表 ----------------
# 项目约定：**运行时不得调用外部地理编码**，所以坐标只能来自本地文件。
#   ① 真实坐标：仓库里既有的一次性高德研究缓存（企业全称命中即用，保留原 precision/status）
#   ② 区域参考：只有注册地址、没有门牌级坐标的客户，按属地**地级市**取本地底图里的
#      地名点近似落位。坐标本身是真实的（OSM/Protomaps 底图里的城市点位），
#      但对企业而言是「区域参考」，不代表门牌级精度 —— 界面上会明确标注。
_ADMIN_SUFFIX = re.compile(
    r"(省|市|自治区|特别行政区|地区|自治州|自治县|盟|区|县|旗|新区|开发区|高新区)$"
)


def load_location_cache():
    """加载本地点位缓存：企业全称 -> 真实坐标。"""
    cache = {}
    for rel in LOCATION_CACHE_FILES:
        path = os.path.join(BASE, rel)
        if not os.path.exists(path):
            continue
        try:
            payload = json.load(open(path, encoding="utf-8"))
        except Exception:
            continue
        records = payload.get("records") or payload.get("candidates") or []
        if isinstance(records, dict):
            records = list(records.values())
        for rec in records:
            if not isinstance(rec, dict):
                continue
            name = norm(rec.get("name"))
            if name and rec.get("lng") and rec.get("lat"):
                cache.setdefault(name, {**rec, "_cacheFile": os.path.basename(rel)})
    return cache


def _load_provinces():
    """加载省界，返回 [(省名, [多边形...], (minLng, minLat, maxLng, maxLat)), ...]。

    每个多边形是 [外环, ...内环]，环是 [[lng, lat], ...]。bbox 用于快速排除。
    """
    if not os.path.exists(PROVINCE_GEOJSON):
        return []
    data = json.load(open(PROVINCE_GEOJSON, encoding="utf-8"))
    out = []
    for feat in data.get("features", []):
        geom = feat.get("geometry") or {}
        name = norm((feat.get("properties") or {}).get("name"))
        if not name:
            continue
        if geom.get("type") == "Polygon":
            polygons = [geom.get("coordinates") or []]
        elif geom.get("type") == "MultiPolygon":
            polygons = geom.get("coordinates") or []
        else:
            continue
        xs = []
        ys = []
        for rings in polygons:
            for ring in rings:
                for pt in ring:
                    xs.append(pt[0])
                    ys.append(pt[1])
        if not xs:
            continue
        out.append((name, polygons, (min(xs), min(ys), max(xs), max(ys))))
    return out


PROVINCES = _load_provinces()


def _ring_contains(ring, lng, lat):
    """射线法：点是否在单个环内。"""
    inside = False
    j = len(ring) - 1
    for i in range(len(ring)):
        xi = ring[i][0]
        yi = ring[i][1]
        xj = ring[j][0]
        yj = ring[j][1]
        if (yi > lat) != (yj > lat) and lng < (xj - xi) * (lat - yi) / (yj - yi) + xi:
            inside = not inside
        j = i
    return inside


def _polygon_contains(rings, lng, lat):
    if not rings or not _ring_contains(rings[0], lng, lat):
        return False
    for hole in rings[1:]:
        if _ring_contains(hole, lng, lat):
            return False
    return True


_PROVINCE_CACHE = {}


def province_of(lng, lat):
    """返回坐标所在的省名；不在任何省界内（含境外）返回 ""。带缓存，避免重复判定。"""
    key = (round(lng, 4), round(lat, 4))
    if key in _PROVINCE_CACHE:
        return _PROVINCE_CACHE[key]
    found = ""
    for name, polygons, bbox in PROVINCES:
        if not (bbox[0] <= key[0] <= bbox[2] and bbox[1] <= key[1] <= bbox[3]):
            continue
        if any(_polygon_contains(rings, lng, lat) for rings in polygons):
            found = name
            break
    _PROVINCE_CACHE[key] = found
    return found


def load_gazetteer():
    """加载地名词表：地名 -> [候选点...]；同时建「去行政后缀」索引。

    同名地名**必须保留全部候选**：既有跨省重名（「朝阳区」「和平区」「铁西区」），
    也有跨国重名（韩国京畿道「广州市」）。原先按名字只留一条，导致中国的广州市
    被韩国的顶掉。消歧交给 gazetteer_point 按省份做。
    """
    path = os.path.join(BASE, GAZETTEER_FILE)
    if not os.path.exists(path):
        return {}, {}
    payload = json.load(open(path, encoding="utf-8"))
    by_name = {}
    by_base = {}
    for rec in payload.get("records", []):
        name = norm(rec.get("name"))
        if not name:
            continue
        by_name.setdefault(name, []).append(rec)
        base = _ADMIN_SUFFIX.sub("", name)
        if base:
            by_base.setdefault(base, []).append(rec)
    return by_name, by_base


LOCATION_CACHE = load_location_cache()
GAZETTEER, GAZETTEER_BASE = load_gazetteer()

# 自治州/地区：底图地名点常以其驻地城市命名，这里给出驻地对照（行政常识，非编造坐标）
PREFECTURE_SEAT = {
    "凉山彝族自治州": "西昌市",
    "湘西土家族苗族自治州": "吉首市",
    "恩施土家族苗族自治州": "恩施市",
    "黔东南苗族侗族自治州": "凯里市",
    "黔南布依族苗族自治州": "都匀市",
    "黔西南布依族苗族自治州": "兴义市",
    "大理白族自治州": "大理市",
    "红河哈尼族彝族自治州": "蒙自市",
    "楚雄彝族自治州": "楚雄市",
    "文山壮族苗族自治州": "文山市",
    "西双版纳傣族自治州": "景洪市",
    "德宏傣族景颇族自治州": "芒市",
    "迪庆藏族自治州": "香格里拉市",
    "怒江傈僳族自治州": "泸水市",
    "临夏回族自治州": "临夏市",
    "甘南藏族自治州": "合作市",
    "海西蒙古族藏族自治州": "德令哈市",
    "海南藏族自治州": "共和县",
    "海北藏族自治州": "海晏县",
    "黄南藏族自治州": "同仁市",
    "果洛藏族自治州": "玛沁县",
    "玉树藏族自治州": "玉树市",
    "伊犁哈萨克自治州": "伊宁市",
    "博尔塔拉蒙古自治州": "博乐市",
    "巴音郭楞蒙古自治州": "库尔勒市",
    "昌吉回族自治州": "昌吉市",
    "克孜勒苏柯尔克孜自治州": "阿图什市",
    "阿克苏地区": "阿克苏市",
    "喀什地区": "喀什市",
    "和田地区": "和田市",
    "塔城地区": "塔城市",
    "阿勒泰地区": "阿勒泰市",
    "兴安盟": "乌兰浩特市",
    "锡林郭勒盟": "锡林浩特市",
    "阿拉善盟": "阿拉善左旗",
    "呼伦贝尔市": "海拉尔区",
    "延边朝鲜族自治州": "延吉市",
    "甘孜藏族自治州": "康定市",
    "阿坝藏族羌族自治州": "马尔康市",
    "大兴安岭地区": "加格达奇区",
}


def _pick_in_province(pool, wanted):
    """从同名候选里挑出落在目标省份内的那条；没有省份信息时退回第一个候选。"""
    if not pool:
        return None
    if not wanted:
        return pool[0]
    for rec in pool:
        if province_of(rec["lng"], rec["lat"]) == wanted:
            return rec
    return None


def city_text(province, city):
    """返回可信的地级市名；若词表显示该市不在标称省份内，判为源数据错误，置空。

    实际案例：源表把「特变电工」的地级市填成「北京市」（注册地址含「北京南路」所致），
    而它实际在新疆昌吉。坐标已按区县纠正为昌吉市，这里再把不属于该省的市名去掉，
    免得地图的省市筛选器里冒出「新疆维吾尔自治区 → 北京市」这种选项。
    词表里查不到该市时无从判断，原样保留（不误伤）。
    """
    name = norm(city)
    prov = norm(province)
    if not name or name == "—" or not prov or name == prov:
        return name
    pool = GAZETTEER.get(name) or GAZETTEER_BASE.get(_ADMIN_SUFFIX.sub("", name))
    if not pool:
        return name
    if any(province_of(rec["lng"], rec["lat"]) == prov for rec in pool):
        return name
    return ""


def gazetteer_point(district, city, province):
    """按属地找地名词表点位：区县 → 地级市 → （自治州/地区）驻地 → 省。

    每一级都要求候选点落在**属地标称的省份**内，这是必须的：词表里存在大量同名
    地名（跨省的「朝阳区」「和平区」，以及跨国的同名点），不校验就会把企业放到
    错误省份，甚至放到国外。
    先试区县还有两个好处：比市名更精确，且能绕开市名层面的同名冲突。
    """
    wanted = norm(province)
    tried = set()
    for raw in (district, city, PREFECTURE_SEAT.get(norm(city)), province):
        name = norm(raw)
        if not name or name == "—" or name in tried:
            continue
        tried.add(name)
        base = _ADMIN_SUFFIX.sub("", name)
        for pool in (GAZETTEER.get(name), GAZETTEER_BASE.get(base)):
            picked = _pick_in_province(pool, wanted)
            if picked:
                return {
                    "lng": picked["lng"],
                    "lat": picked["lat"],
                    "matchedName": picked["name"],
                }
    return None


def resolve_location(name, district, city, province):
    """返回 (location, locationStatus, locationPrecision, locationSource, matchedName)。"""
    wanted = norm(province)
    cached = LOCATION_CACHE.get(norm(name))
    if cached:
        lng = float(cached["lng"])
        lat = float(cached["lat"])
        # 缓存里混着低置信度匹配（正是本项目历史上坐标错位的来源）。用省界兜底：
        # 坐标不在标称省份内就丢弃，改用区域参考落点 —— 宁可粗，不能错。
        if not wanted or province_of(lng, lat) == wanted:
            return (
                {
                    "lng": round(lng, 6),
                    "lat": round(lat, 6),
                    "coordinateSystem": "WGS84/CGCS2000-compatible",
                    "precision": cached.get("precision") or "候选",
                },
                cached.get("status") or "候选",
                cached.get("precision") or "候选",
                f"本地点位缓存（{cached.get('_cacheFile')}）",
                cached.get("matchedAddress") or "",
            )
    point = gazetteer_point(district, city, province)
    if point:
        return (
            {
                "lng": point["lng"],
                "lat": point["lat"],
                "coordinateSystem": "WGS84/CGCS2000-compatible",
                "precision": "区域参考",
            },
            "区域参考",
            "区域参考",
            f"本地底图地名点（{point['matchedName']}）",
            point["matchedName"],
        )
    return (None, "待补充", "", "", "")


def main():
    # ---------- Sheet1：节点框架 ----------
    wb = openpyxl.load_workbook(SRC_CHAIN, data_only=True)
    ws1 = wb["Sheet1"]
    node_rows = []
    for r in range(2, ws1.max_row + 1):
        pos, l1, l2, order = (ws1.cell(r, c).value for c in range(1, 5))
        if not norm(l1):
            continue
        node_rows.append(
            {
                "position": norm(pos),
                "level1": norm(l1),
                "level2": norm(l2),
                "order": int(order) if order is not None else 0,
            }
        )
    node_rows.sort(key=lambda x: x["order"])

    # ---------- Sheet3：企业 ----------
    ws3 = wb["Sheet3"]
    header = {ws3.cell(1, c).value: c for c in range(1, ws3.max_column + 1)}
    rows = []
    for r in range(2, ws3.max_row + 1):
        rec = {k: ws3.cell(r, v).value for k, v in header.items()}
        if not norm(rec.get("上市企业")):
            continue
        rows.append(
            {
                "name": norm(rec.get("上市企业")),
                "fullName": norm(rec.get("企业名称")),
                "code": norm(rec.get("股票代码")),
                "board": norm(rec.get("板块")),
                "swSector": norm(rec.get("申万一级")),
                "address": norm(rec.get("地址")),
                "products": norm(rec.get("主要储能产品/业务")),
                "evidence": norm(rec.get("年报/行业信息依据")),
                "grade": norm(rec.get("储能相关度档位")),
                "fillMark": norm(rec.get("补入标记")),
                "level1": norm(rec.get("一级环节")),
                "level2": norm(rec.get("二级环节")),
            }
        )

    # ---------- 头部标记表 ----------
    wb_h = openpyxl.load_workbook(SRC_HEADS, data_only=True)
    ws_h = wb_h["01_分环节头部企业"]
    hh = {ws_h.cell(1, c).value: c for c in range(1, ws_h.max_column + 1)}
    head_rows = []
    for r in range(2, ws_h.max_row + 1):
        if not norm(ws_h.cell(r, 1).value):
            continue
        head_rows.append({k: ws_h.cell(r, v).value for k, v in hh.items()})

    # name -> tier；name -> [细分环节]
    tier_of = {}
    segments_of = {}
    conflicts = []
    for rec in head_rows:
        segment = norm(rec.get("环节"))
        for raw_name, raw_tier in HEAD_PATTERN.findall(norm(rec.get("头部企业（龙头/骨干）"))):
            name = raw_name.strip()
            tier = "leader" if raw_tier == "龙头" else "core"
            if name in tier_of and tier_of[name] != tier:
                conflicts.append({"name": name, "tiers": [tier_of[name], tier]})
            tier_of[name] = tier
            segments_of.setdefault(name, [])
            if segment not in segments_of[name]:
                segments_of[name].append(segment)

    # 应用别名，把头部名单落到企业表的名字上
    head_lookup = {}
    alias_hits = []
    for name, tier in tier_of.items():
        target = NAME_ALIASES.get(name, name)
        if target != name:
            alias_hits.append({"from": name, "to": target})
        head_lookup[target] = tier

    # ---------- 关系档案（股权穿透 / 关键人 / 供销 / 偏好） ----------
    # 六个 sheet 均以「上市简称」为键，一企一行；本项目按简称落到企业表，
    # 并沿用 NAME_ALIASES（五矿新能→长远锂科、青鸟智控→青鸟消防）。
    wb_rel = openpyxl.load_workbook(SRC_REL, data_only=True)

    def sheet_rows(sheet_name):
        ws = wb_rel[sheet_name]
        hdr = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        out = collections.OrderedDict()
        for r in range(2, ws.max_row + 1):
            rec = {k: ws.cell(r, v).value for k, v in hdr.items()}
            key = norm(rec.get("上市简称"))
            if key:
                out[key] = rec
        return out

    rel_summary = sheet_rows("01_关系档案汇总")
    rel_equity = sheet_rows("02_股权穿透")
    rel_people = sheet_rows("03_关键人")
    rel_supply = sheet_rows("04_供销关系")
    rel_prefs = sheet_rows("05_关键人偏好")
    rel_territory = sheet_rows("07_属地信息") if "07_属地信息" in wb_rel.sheetnames else {}

    # ---------- 属地（07_属地信息）：省 / 市 / 区县 / 街道 / 园区 / 注册地址 ----------
    # 以工商登记「注册地址」为准解析五级属地。注意两类特殊行：
    #   ① 区县为空 = 注册地址是「园区/开发区/道路」型（源数据本身未写区县），不是解析失败
    #   ② 离岸注册（华润微=开曼、华虹公司=中国香港）—— 属地在境外，但境内有运营主体
    territory_by_short = {}
    for short, rec in rel_territory.items():
        target = NAME_ALIASES.get(short, short)
        territory_by_short[target] = {
            "province": norm(rec.get("省/自治区/直辖市")),
            "city": norm(rec.get("地级市")),
            "district": norm(rec.get("区县")),
            "street": norm(rec.get("所在街道/乡镇")),
            "park": norm(rec.get("所在园区/开发区")),
            "address": norm(rec.get("注册地址")),
            "note": norm(rec.get("备注")),
            "source": norm(rec.get("数据来源")),
        }

    def territory_for(short):
        return territory_by_short.get(short, {})

    # ---------- 全量客户清单：节点关系（官方编码）+ 非 237 客户 ----------
    wb_all = openpyxl.load_workbook(SRC_ALL, data_only=True)

    ws_nodes = wb_all["节点关系"]
    hdr_nodes = {ws_nodes.cell(1, c).value: c for c in range(1, ws_nodes.max_column + 1)}
    official_code = {}  # (一级环节, 二级环节) -> U1…D6
    official_order = {}
    node_position = {}
    for r in range(2, ws_nodes.max_row + 1):
        l1 = norm(ws_nodes.cell(r, hdr_nodes["一级环节"]).value)
        l2 = norm(ws_nodes.cell(r, hdr_nodes["二级环节"]).value)
        if not l1 or not l2:
            continue
        official_code[(l1, l2)] = norm(ws_nodes.cell(r, hdr_nodes["二级环节编码"]).value)
        official_order[(l1, l2)] = ws_nodes.cell(r, hdr_nodes["顺序"]).value
        node_position[(l1, l2)] = norm(ws_nodes.cell(r, hdr_nodes["位置"]).value)

    # 与主表（Sheet1）的 24 个节点逐一比对，不一致就明确报出来，不静默取一个
    node_mismatch = sorted(
        {p for p in official_code if p not in {(n["level1"], n["level2"]) for n in node_rows}}
        | {(n["level1"], n["level2"]) for n in node_rows if (n["level1"], n["level2"]) not in official_code}
    )

    ws_other = wb_all["非237客户范围"]
    hdr_other = {ws_other.cell(1, c).value: c for c in range(1, ws_other.max_column + 1)}
    other_rows = []
    for r in range(2, ws_other.max_row + 1):
        rec = {k: ws_other.cell(r, v).value for k, v in hdr_other.items()}
        name = norm(rec.get("企业名称"))
        if not name:
            continue
        codes = [
            c.strip()
            for c in re.split(r"[;；,，、]", norm(rec.get("切法B节点")))
            if c.strip()
        ]
        other_rows.append(
            {
                "name": name,
                "creditCode": norm(rec.get("统一社会信用代码")),
                "codes": codes,
                "nodeDetail": norm(rec.get("节点详情")),
                "address": norm(rec.get("地址")),
                "mainBusiness": norm(rec.get("主营业务摘要")),
                "source": norm(rec.get("来源")),
                "verifyStatus": norm(rec.get("核验状态")),
                "province": norm(rec.get("属地(省)")),
                "city": norm(rec.get("属地(市)")),
                "district": norm(rec.get("属地(区县)")),
            }
        )

    # 编码 -> (一级环节, 二级环节)
    code_to_node = {v: k for k, v in official_code.items() if v}

    # 关系档案完整度口径（来自 06_数据说明与完整度，用于展示覆盖情况）
    # ⚠️ 该 sheet 用「行首缩进」标记子项，不能用 norm()（会把缩进 strip 掉）。
    ws_note = wb_rel["06_数据说明与完整度"]
    rel_coverage = []
    for r in range(1, ws_note.max_row + 1):
        raw_label = ws_note.cell(r, 1).value
        value = norm(ws_note.cell(r, 2).value)
        if not isinstance(raw_label, str) or not raw_label.startswith("  ") or not value:
            continue
        key = raw_label.strip().split(" ")[0]
        parts = dict(
            (seg.strip().split(" ")[0], int(re.sub(r"\D", "", seg) or 0))
            for seg in value.split("/")
        )
        rel_coverage.append(
            {
                "key": key,
                "full": parts.get("完整", 0),
                "partial": parts.get("部分", 0),
                "missing": parts.get("缺失", 0),
            }
        )

    company_short_names = {r["name"] for r in rows}

    rel_profiles = {}
    rel_alias_hits = []
    for short, rec in rel_summary.items():
        target = NAME_ALIASES.get(short, short)
        if target != short:
            rel_alias_hits.append({"from": short, "to": target})
        if target not in company_short_names:
            continue

        def cell(sheet, field):
            return norm((sheet.get(short) or {}).get(field))

        rel_profiles[target] = {
            "nature": norm(rec.get("企业性质")),
            "industry": norm(rec.get("所属行业")),
            "controller": norm(rec.get("实际控制人")),
            "controllerType": norm(rec.get("实控人性质")),
            "isStateOwned": norm(rec.get("是否国资控股")) == "是",
            "topShareholder": norm(rec.get("第一大股东及持股")),
            "controlPath": cell(rel_equity, "控制路径/持股比例"),
            "top10": cell(rel_equity, "前十大股东摘要"),
            "chairman": norm(rec.get("董事长")) or cell(rel_people, "董事长"),
            "legalPerson": norm(rec.get("法定代表人")),
            "generalManager": norm(rec.get("总经理/总裁")),
            "founderBackground": cell(rel_people, "创始人/实控人背景"),
            "peopleNote": cell(rel_people, "关键人简介"),
            "customers": cell(rel_supply, "前五大客户"),
            "suppliers": cell(rel_supply, "前五大供应商"),
            "chainNote": cell(rel_supply, "主要上下游配套"),
            "supplyDisclosed": cell(rel_supply, "是否披露具体名称"),
            "prefStrategy": cell(rel_prefs, "战略关注方向"),
            "prefEvents": cell(rel_prefs, "公开活动/参会会展"),
            "prefAssociations": cell(rel_prefs, "行业协会任职"),
            "prefSpeeches": cell(rel_prefs, "公开发言主题"),
            "prefNote": cell(rel_prefs, "备注"),
            "completeness": {
                "equity": norm(rec.get("股权完整度")),
                "keyPersons": norm(rec.get("关键人完整度")),
                "supplyChain": norm(rec.get("供销完整度")),
                "preferences": norm(rec.get("偏好完整度")),
            },
        }

    rel_missing = sorted(company_short_names - set(rel_profiles))
    rel_unmatched = sorted(set(rel_summary) - set(NAME_ALIASES) - company_short_names)

    # ---------- 从关系档案派生「链内关系边」 ----------
    # 方向约定与既有分片一致：from 供应 to（relationType=供应）。
    # 若 B 出现在 A 的「前五大客户」→ A 卖给 B；若出现在「前五大供应商」→ B 卖给 A。
    # 只在两端都是本项目 237 家上市企业时保留，避免把外部公司拉进图谱。
    short_to_full = {r["name"]: (r["fullName"] or r["name"]) for r in rows}
    name_by_len = sorted(company_short_names, key=len, reverse=True)

    def names_in_text(text, exclude):
        """在自由文本里找出属于本链的企业简称（长名优先，避免短名误命中）。"""
        # 先剔除括号内的口径/比例说明，再匹配
        flat = re.sub(r"[（(][^）)]*[）)]", " ", text or "")
        found = []
        for name in name_by_len:
            if name == exclude or name not in flat:
                continue
            if any(name != other and name in other and other in found for other in found):
                continue
            found.append(name)
        return found

    edge_records = []
    for short, rec in rel_supply.items():
        a = NAME_ALIASES.get(short, short)
        if a not in company_short_names:
            continue
        for field, kind in (("前五大客户", "客户"), ("前五大供应商", "供应商")):
            for other in names_in_text(norm(rec.get(field)), a):
                edge_records.append((a, other, kind))

    # 去重
    dedup = collections.OrderedDict()
    for a, b, kind in edge_records:
        dedup.setdefault((a, b, kind), True)
    edge_records = [(a, b, kind) for (a, b, kind) in dedup]

    # 股权关系：实控人或第一大股东本身就是本链企业
    equity_edges = []
    for short, rec in rel_equity.items():
        a = NAME_ALIASES.get(short, short)
        if a not in company_short_names:
            continue
        blob = f"{norm(rec.get('实际控制人'))} {norm(rec.get('第一大股东及持股比例'))}"
        for holder in names_in_text(blob, a):
            equity_edges.append((holder, a))

    relations = []
    for i, (a, b, kind) in enumerate(edge_records, start=1):
        frm, to = (a, b) if kind == "客户" else (b, a)
        relations.append(
            {
                "id": f"relation-es-{i:04d}",
                "from": short_to_full[frm],
                "to": short_to_full[to],
                "relationType": "供应",
                "depth": 1,
                "source": "上市企业关系档案（前五大客户/供应商）",
                "sourceFromProfile": frm,
                "sourceToProfile": to,
                "basis": f"{a} 前五大{'客户' if kind == '客户' else '供应商'}中列示 {b}",
            }
        )
    base = len(relations)
    for j, (holder, target) in enumerate(equity_edges, start=1):
        relations.append(
            {
                "id": f"relation-es-e{j:03d}",
                "from": short_to_full[holder],
                "to": short_to_full[target],
                "relationType": "股权",
                "depth": 1,
                "source": "上市企业关系档案（股权穿透）",
                "sourceFromProfile": holder,
                "sourceToProfile": target,
                "basis": f"{holder} 为 {target} 的控股股东 / 实际控制人",
            }
        )

    relation_count_by_short = collections.Counter()
    for rel in relations:
        relation_count_by_short[rel["sourceFromProfile"]] += 1
        relation_count_by_short[rel["sourceToProfile"]] += 1

    relations_stats = {
        "total": len(relations),
        "supply": base,
        "equity": len(relations) - base,
        "companies": len(relation_count_by_short),
        "profiles": len(rel_profiles),
    }

    # ---------- 链外被点名主体（补口径，237 统计不受影响） ----------
    # ⚠️ 237 家是「上市企业」样本，产业链本身不只含上市公司。上市公司的年报/公开披露里
    #    大量点名非上市主体：上游如 天齐锂业的供应商「河北吉诚」（非上市，电池级碳酸锂 /
    #    氢氧化锂，属本链「资源与材料」环节），下游如 国家电网、广汽埃安、因湃电池。
    #    上面成边规则要求「两端都在 237 内」，会把这些整条丢掉，因此在数据里单独留一层，
    #    只做事后展示，不参与 237 的头部统计与关系条数。
    external_raw = {}  # name -> {upstream:[(owner, ctx)], downstream:[...]}
    # 237 家的简称 + 全称 + 别名，用来把「本链企业」排除出链外名单
    known_names = set(company_short_names) | {
        r["fullName"] for r in rows if r.get("fullName")
    }
    known_names |= set(NAME_ALIASES)
    for short, rec in rel_supply.items():
        owner = NAME_ALIASES.get(short, short)
        if owner not in company_short_names:
            continue
        for field, role in EXTERNAL_FIELDS:
            raw = norm(rec.get(field))
            if not raw:
                continue
            text = re.sub(r"[（(][^）)]*[）)]", " ", raw)
            text = re.sub(r"^前五[大名称]{1,2}(客户|供应商)?[:：]?", " ", text)
            for tok in re.split(r"[、；;，,／/\n]+", text):
                tok = re.sub(r"^[^：:]{0,10}[:：]", "", tok).strip(" -—·:：.*")
                tok = EXTERNAL_TRAIL.sub("", tok)
                # 本链 237 家自己不算链外
                if any(k and (k in tok or tok in k) for k in known_names):
                    continue
                tok = EXTERNAL_NORMALIZE.get(tok, tok)
                if not tok or tok in EXTERNAL_DROP:
                    continue
                if not looks_like_external_entity(tok):
                    continue
                bucket = external_raw.setdefault(tok, {"upstream": [], "downstream": []})
                i = text.find(tok)
                ctx = re.sub(r"\s+", " ", text[max(0, i - 30) : i + 40]).strip()
                slot = bucket["upstream"] if role == "upstream" else bucket["downstream"]
                slot.append((owner, ctx))

    external_mentions = []
    for name, bucket in external_raw.items():
        up, down = bucket["upstream"], bucket["downstream"]
        if len(up) + len(down) < EXTERNAL_MIN_COUNT:
            continue
        seen_owner = set()
        mentioned_by = []
        for owner, ctx in up + down:
            if owner in seen_owner:
                continue
            seen_owner.add(owner)
            mentioned_by.append(
                {
                    "company": owner,
                    "role": "上游" if (owner, ctx) in up else "下游",
                    "context": ctx,
                }
            )
        external_mentions.append(
            {
                "name": name,
                "count": len(up) + len(down),
                "asUpstream": len(up),
                "asDownstream": len(down),
                "mentionedBy": mentioned_by[:EXTERNAL_MAX_SOURCES],
            }
        )
    external_mentions.sort(key=lambda x: (-x["count"], -x["asUpstream"], x["name"]))

    external_stats = {
        "total": len(external_mentions),
        # 与界面同一口径：主要出现在「供应商栏」即算偏上游，否则算偏下游
        "upstream": sum(
            1
            for e in external_mentions
            if e["asUpstream"] > 0 and e["asUpstream"] >= e["asDownstream"]
        ),
        "downstream": sum(
            1 for e in external_mentions if e["asDownstream"] > e["asUpstream"]
        ),
        "minCount": EXTERNAL_MIN_COUNT,
    }

    # ---------- 属地分布（来自 07_属地信息） ----------
    terr_rows = list(territory_by_short.values())
    province_count = collections.Counter(t["province"] for t in terr_rows if t["province"])
    city_count = collections.Counter(
        t["city"] for t in terr_rows if t["city"] and t["city"] != "—"
    )
    # 离岸注册主体（属地在境外，境内另有运营主体）—— 单独列出，别混进国内分布
    offshore = [
        {"name": short, "province": t["province"], "note": t["note"]}
        for short, t in territory_by_short.items()
        if t["province"] and ("香港" in t["province"] or "境外" in t["province"])
    ]
    territory_stats = {
        "total": len(terr_rows),
        "provinces": len(province_count),
        "cities": len(city_count),
        "withDistrict": sum(1 for t in terr_rows if t["district"]),
        "withStreet": sum(1 for t in terr_rows if t["street"]),
        "withPark": sum(1 for t in terr_rows if t["park"]),
        "offshore": len(offshore),
        "topProvinces": [{"k": k, "v": v} for k, v in province_count.most_common(12)],
        "topCities": [{"k": k, "v": v} for k, v in city_count.most_common(12)],
        "offshoreList": offshore,
        # 区县为空 = 注册地址是「园区/开发区/道路」型，源数据本身未写区县，不是解析失败
        "missingDistrict": [t["address"] for t in terr_rows if not t["district"]],
    }

    # ---------- 组装节点与企业的对应关系 ----------
    level1_order = []
    for n in node_rows:
        if n["level1"] not in level1_order:
            level1_order.append(n["level1"])

    node_companies = collections.OrderedDict(
        ((n["level1"], n["level2"]), []) for n in node_rows
    )
    unmatched_rows = collections.Counter()
    for rec in rows:
        key = (rec["level1"], rec["level2"])
        if key in node_companies:
            node_companies[key].append(rec)
        else:
            unmatched_rows[f"{rec['level1']} / {rec['level2']}"] += 1

    # 企业级 tier（按简称）
    company_names = {r["name"] for r in rows}
    tier_by_company = {}
    for name in company_names:
        tier_by_company[name] = head_lookup.get(name, "tail")
    head_names_not_found = sorted(set(head_lookup) - company_names)

    def company_payload(rec):
        tier = tier_by_company.get(rec["name"], "tail")
        # 注意：节点级企业只放轻量字段。完整关系档案（profile）只在顶层 companies[] 里存一份，
        # 否则同一企业出现在多个二级环节时会把整份档案重复写入，文件体积翻倍。
        return {
            "name": rec["name"],
            "code": rec["code"],
            "board": rec["board"],
            "swSector": rec["swSector"],
            "products": rec["products"],
            "evidence": rec["evidence"],
            "grade": rec["grade"],
            "fillMark": rec["fillMark"],
            "tier": tier,
            "tierLabel": TIER_LABEL[tier],
            "tierSegments": segments_of.get(rec["name"], []),
            "relationCount": relation_count_by_short.get(rec["name"], 0),
            # 占位评分明细：step3 用它做「配置权重 → 重新评估分档」
            "scoreDetail": build_score_detail(rec["name"], tier),
            "sourceTier": tier,
        }

    code_of = {l1: f"{i + 1:02d}" for i, l1 in enumerate(level1_order)}

    # ---------- 非 237 客户：按节点编码归位 + 解析落点 ----------
    other_by_node = collections.OrderedDict()
    other_customers = []
    other_no_node = []
    other_region_ref = 0
    other_real = 0
    other_unplaced = []
    for rec in other_rows:
        location, loc_status, loc_precision, loc_source, matched = resolve_location(
            rec["name"], rec["district"], rec["city"], rec["province"]
        )
        if loc_status == "区域参考":
            other_region_ref += 1
        elif location:
            other_real += 1
        else:
            other_unplaced.append(rec["name"])
        primary = rec["codes"][0] if rec["codes"] else ""
        item = {
            "name": rec["name"],
            "creditCode": rec["creditCode"],
            "codes": rec["codes"],
            "nodeDetail": rec["nodeDetail"],
            "primaryCode": primary,
            "address": rec["address"],
            "mainBusiness": rec["mainBusiness"],
            "source": rec["source"],
            "verifyStatus": rec["verifyStatus"],
            "province": rec["province"],
            "city": rec["city"],
            "district": rec["district"],
            "location": location,
            "locationStatus": loc_status,
            "locationPrecision": loc_precision,
            "locationSource": loc_source,
            "locationMatchedName": matched,
        }
        other_customers.append(item)
        if not primary:
            other_no_node.append(rec["name"])
            continue
        for code in rec["codes"]:
            other_by_node.setdefault(code, []).append(item)

    def other_payload(item):
        """节点内的轻量客户条目（位置与明细在顶层 otherCustomers 已有一份）。"""
        return {
            "name": item["name"],
            "creditCode": item["creditCode"],
            "primaryCode": item["primaryCode"],
            "codes": item["codes"],
            "province": item["province"],
            "city": item["city"],
            "district": item["district"],
            "source": item["source"],
            "verifyStatus": item["verifyStatus"],
            "locationStatus": item["locationStatus"],
        }

    customer_stats = {
        # 此处 companies[] 尚未组装，237 直接用 company_short_names 的口径
        "listed": len(company_short_names),
        "other": len(other_customers),
        "total": len(company_short_names) + len(other_customers),
        "otherListed": other_real + other_region_ref,
        "otherRealLocation": other_real,
        "otherRegionReference": other_region_ref,
        "otherUnplaced": len(other_unplaced),
        "otherBySource": [{"k": k, "v": v} for k, v in
                          collections.Counter(c["source"] for c in other_customers).most_common()],
        "otherByVerify": [{"k": k, "v": v} for k, v in
                          collections.Counter(c["verifyStatus"] for c in other_customers).most_common()],
        "officialCodeMismatch": node_mismatch,
    }

    # ---------- 全量客户属地分布（237 上市 + 1740 非上市） ----------
    all_territory = [
        {
            "province": t["province"],
            "city": t["city"],
            "district": t["district"],
            "listed": True,
        }
        for t in territory_by_short.values()
    ] + [
        {
            "province": c["province"],
            "city": c["city"],
            "district": c["district"] if c["district"] != "—" else "",
            "listed": False,
        }
        for c in other_customers
    ]
    list_prov = collections.Counter(
        t["province"] for t in all_territory if t["listed"] and t["province"]
    )
    other_prov = collections.Counter(
        t["province"] for t in all_territory if not t["listed"] and t["province"]
    )
    list_city = collections.Counter(
        t["city"] for t in all_territory if t["listed"] and t["city"] and t["city"] != "—"
    )
    other_city = collections.Counter(
        t["city"] for t in all_territory if not t["listed"] and t["city"] and t["city"] != "—"
    )
    prov_keys = [k for k, _ in (list_prov + other_prov).most_common(14)]
    city_keys = [k for k, _ in (list_city + other_city).most_common(14)]
    customer_territory = {
        "listed": len(territory_by_short),
        "other": len(other_customers),
        "total": len(all_territory),
        "provinces": len(set(list_prov) | set(other_prov)),
        "cities": len(set(list_city) | set(other_city)),
        "topProvinces": [
            {"k": k, "listed": list_prov.get(k, 0), "other": other_prov.get(k, 0)}
            for k in prov_keys
        ],
        "topCities": [
            {"k": k, "listed": list_city.get(k, 0), "other": other_city.get(k, 0)}
            for k in city_keys
        ],
    }


    # ---------- levels（7 个一级环节，含 24 个二级环节） ----------
    levels = []
    for idx, l1 in enumerate(level1_order, start=1):
        children = []
        for n in node_rows:
            if n["level1"] != l1:
                continue
            pool = node_companies[(n["level1"], n["level2"])]
            pool_sorted = sorted(
                pool,
                key=lambda c: (
                    {"leader": 0, "core": 1, "tail": 2}[tier_by_company.get(c["name"], "tail")],
                    GRADE_ORDER.index(c["grade"]) if c["grade"] in GRADE_ORDER else 99,
                    c["name"],
                ),
            )
            grade_count = collections.Counter(c["grade"] for c in pool_sorted)
            tier_count = collections.Counter(
                tier_by_company.get(c["name"], "tail") for c in pool_sorted
            )

            # 该二级环节下的细分环节头部分组（来自头部标记表的「环节」列）
            head_segments = []
            for rec in head_rows:
                if norm(rec.get("一级环节")) != n["level1"]:
                    continue
                if norm(rec.get("二级环节")) != n["level2"]:
                    continue
                members = []
                for raw_name, raw_tier in HEAD_PATTERN.findall(
                    norm(rec.get("头部企业（龙头/骨干）"))
                ):
                    short = raw_name.strip()
                    target = NAME_ALIASES.get(short, short)
                    members.append(
                        {
                            "name": target,
                            "tier": "leader" if raw_tier == "龙头" else "core",
                            "tierLabel": raw_tier,
                        }
                    )
                head_segments.append(
                    {
                        "segment": norm(rec.get("环节")),
                        "declaredTotal": int(rec.get("总数") or 0),
                        "declaredHead": int(rec.get("头部企业数") or 0),
                        "leaderCount": sum(1 for m in members if m["tier"] == "leader"),
                        "coreCount": sum(1 for m in members if m["tier"] == "core"),
                        "members": members,
                    }
                )

            children.append(
                {
                    "order": n["order"],
                    "nodeKey": f"{n['order']:02d}-{n['level1']}-{n['level2']}",
                    # 官方二级环节编码（U1…D6），来自全量客户清单的「节点关系」sheet
                    "code2": official_code.get((n["level1"], n["level2"]), ""),
                    "label": n["level2"],
                    "note": LEVEL2_NOTE.get(n["level2"], ""),
                    "companyCount": len(pool_sorted),
                    "grades": [
                        {"grade": g, "count": grade_count[g]}
                        for g in GRADE_ORDER
                        if grade_count.get(g)
                    ],
                    "leaderCount": tier_count.get("leader", 0),
                    "coreCount": tier_count.get("core", 0),
                    "tailCount": tier_count.get("tail", 0),
                    # 非 237 的本链客户（按同一节点编码归位）
                    "otherCount": len(other_by_node.get(
                        official_code.get((n["level1"], n["level2"]), ""), []
                    )),
                    "others": [
                        other_payload(c)
                        for c in other_by_node.get(
                            official_code.get((n["level1"], n["level2"]), ""), []
                        )
                    ],
                    "companies": [company_payload(c) for c in pool_sorted],
                    "heads": head_segments,
                }
            )

        levels.append(
            {
                "code": f"{idx:02d}",
                "label": l1,
                "position": sorted({n["position"] for n in node_rows if n["level1"] == l1})[0],
                "note": LEVEL1_NOTE.get(l1, ""),
                "order": min(c["order"] for c in children),
                "childCount": len(children),
                "companyCount": len(
                    {c["name"] for ch in children for c in ch["companies"]}
                ),
                "companyRowCount": sum(c["companyCount"] for c in children),
                "leaderCount": sum(c["leaderCount"] for c in children),
                "coreCount": sum(c["coreCount"] for c in children),
                "tailCount": len(
                    {
                        c["name"]
                        for ch in children
                        for c in ch["companies"]
                        if c["tier"] == "tail"
                    }
                ),
                "children": children,
            }
        )

    # ---------- 全局去重企业 ----------
    by_company = collections.OrderedDict()
    for rec in rows:
        entry = by_company.setdefault(
            rec["name"],
            {
                "name": rec["name"],
                "fullName": rec["fullName"],
                "code": rec["code"],
                "board": rec["board"],
                "swSector": rec["swSector"],
                "address": rec["address"],
                "products": rec["products"],
                "evidence": rec["evidence"],
                "grade": rec["grade"],
                "fillMark": rec["fillMark"],
                "tier": tier_by_company.get(rec["name"], "tail"),
                "tierLabel": TIER_LABEL[tier_by_company.get(rec["name"], "tail")],
                "tierSegments": segments_of.get(rec["name"], []),
                "relationCount": relation_count_by_short.get(rec["name"], 0),
                "profile": rel_profiles.get(rec["name"]),
                # 属地（省/市/区县/街道/园区 + 注册地址）来自 07_属地信息
                "territory": territory_for(rec["name"]) or None,
                "level1": [],
                "level2": [],
                "nodes": [],
            },
        )
        if rec["level1"] not in entry["level1"]:
            entry["level1"].append(rec["level1"])
        if rec["level2"] not in entry["level2"]:
            entry["level2"].append(rec["level2"])
        entry["nodes"].append(f"{rec['level1']} / {rec['level2']}")
    companies = list(by_company.values())

    tier_total = collections.Counter(c["tier"] for c in companies)

    # ---------- positions ----------
    positions = []
    pos_order = []
    for n in node_rows:
        if n["position"] not in pos_order:
            pos_order.append(n["position"])
    for p in pos_order:
        inner = [lv for lv in levels if lv["position"] == p]
        positions.append(
            {
                "id": p,
                "label": p,
                "note": POSITION_NOTE.get(p, ""),
                "level1Count": len(inner),
                "level2Count": sum(lv["childCount"] for lv in inner),
                "companyCount": len(
                    {c["name"] for lv in inner for ch in lv["children"] for c in ch["companies"]}
                ),
                "span": {
                    "start": min(lv["order"] for lv in inner),
                    "end": max(max(ch["order"] for ch in lv["children"]) for lv in inner),
                },
            }
        )

    # ---------- 统计 ----------
    grade_count = collections.Counter(r["grade"] for r in rows)
    fill_count = collections.Counter(r["fillMark"] for r in rows)
    board_count = collections.Counter(r["board"] for r in rows)
    sw_count = collections.Counter(r["swSector"] for r in rows)
    level2_total = sum(lv["childCount"] for lv in levels)
    empty_nodes = [
        f"{lv['label']} / {ch['label']}"
        for lv in levels
        for ch in lv["children"]
        if ch["companyCount"] == 0
    ]
    nodes_without_heads = [
        f"{lv['label']} / {ch['label']}"
        for lv in levels
        for ch in lv["children"]
        if not ch["heads"] and ch["companyCount"] > 0
    ]

    stats = {
        "level1": len(levels),
        "level2": level2_total,
        "rows": len(rows),
        "companies": len(companies),
        "leader": tier_total.get("leader", 0),
        "core": tier_total.get("core", 0),
        "tail": tier_total.get("tail", 0),
        "coreMain": grade_count.get("核心主业", 0),
        "coreGrade": grade_count.get("核心", 0),
        "important": grade_count.get("相关重要", 0),
        "edge": grade_count.get("边缘观察", 0),
        "related": grade_count.get("关联", 0),
        "filled": fill_count.get("补入", 0),
        "kept": fill_count.get("保留", 0),
        "pool": fill_count.get("观察池", 0),
        "emptyNodes": len(empty_nodes),
        "nodesWithoutHeads": len(nodes_without_heads),
        "relations": relations_stats["total"],
        "relationSupply": relations_stats["supply"],
        "relationEquity": relations_stats["equity"],
        "relationCompanies": relations_stats["companies"],
        "relationProfiles": relations_stats["profiles"],
        "externalMentions": external_stats["total"],
        "externalUpstream": external_stats["upstream"],
        "externalDownstream": external_stats["downstream"],
    }

    # ---------- 四步 ----------
    steps = [
        {
            "id": "nodes",
            "index": 1,
            "title": "形成环节骨架",
            "note": "先定一级、二级环节，把链的骨架搭起来",
            "detail": (
                "依据权威资料与研报、工信部与行业协会口径确定一级/二级环节。"
                f"{INDUSTRY_LABEL}当前为 {len(levels)} 个一级环节、{level2_total} 个二级环节，"
                "自上游向下游并排展开。"
            ),
            "status": "ready",
            "stats": {"level1": len(levels), "level2": level2_total},
        },
        {
            "id": "companies",
            "index": 2,
            "title": "上市企业入链",
            "note": "把候选上市企业按二级环节归位，标注储能相关度",
            "detail": (
                "来源包括上市公司年报与公告、行业协会与工信部信息、发改委与人民政府重点项目名单、"
                "舆情与技术线索。同一企业可归属多个二级环节。"
                f"本步口径为「上市企业」，共 {len(companies)} 家；另有来自 CNESA / CIAPS 的 "
                f"{len(other_customers)} 家**非上市**客户按同一套节点编码归位，"
                f"合计 {len(companies) + len(other_customers)} 家全量客户可见于地图「全部企业」范围。"
            ),
            "status": "ready",
            "sources": [
                {"label": "上市公司年报与公告", "note": "企业主营与储能业务披露"},
                {"label": "行业协会 · 工信部", "note": "中关村储能联盟、电源行业协会"},
                {"label": "发改委 · 人民政府", "note": "省重大项目名单", "highlight": True},
                {"label": "舆情与技术线索", "note": "补充其他企业"},
                {"label": "权威资料与研究", "note": "券商与行业研究报告"},
            ],
            "stats": {
                "companies": len(companies),
                "rows": len(rows),
                "other": len(other_customers),
                "total": len(companies) + len(other_customers),
                "coreMain": stats["coreMain"],
                "core": stats["coreGrade"],
                "important": stats["important"],
                "edge": stats["edge"],
                "related": stats["related"],
            },
        },
        {
            "id": "heads",
            "index": 3,
            "title": "头部企业打标",
            "note": "给入链企业打上龙头 / 骨干 / 长尾标签",
            "detail": (
                f"按综合得分分档：国内份额与排名 40 + 技术壁垒 25 + 规模与盈利 20 + 客户结构与成长 15。"
                f"得分 ≥80 且过硬门槛为龙头，60–79 为骨干，<60 为长尾。"
                f"储能产业 {len(companies)} 家上市企业已全部评档。"
            ),
            "status": "ready",
            "rule": HEAD_RULE,
            "stats": {
                "total": len(companies),
                "leader": stats["leader"],
                "core": stats["core"],
                "tail": stats["tail"],
            },
        },
        {
            "id": "relations",
            "index": 4,
            "title": "关键关系展示",
            "note": "展示上市企业之间的供销、股权关系",
            "detail": (
                "以关系档案中的「前五大客户 / 供应商」具名披露与「股权穿透」为依据。"
                "关系边只保留两端都在本链 237 家上市企业之内的关系，"
                f"当前识别 {relations_stats['total']} 条"
                f"（供销 {relations_stats['supply']} · 股权 {relations_stats['equity']}），"
                f"覆盖 {relations_stats['companies']} 家企业；"
                f"被点名但不在 237 内的链外主体另有 {external_stats['total']} 个，单列在「链外关键配套」。"
            ),
            "status": "ready",
            "dimensions": [
                {
                    "key": "equity",
                    "label": "股权穿透",
                    "note": "实际控制人 · 控制路径 · 前十大股东",
                },
                {"key": "keyPersons", "label": "关键人", "note": "董事长 · 法定代表人 · 总经理"},
                {"key": "supplyChain", "label": "供销关系", "note": "前五大客户 / 供应商"},
                {"key": "preferences", "label": "关键人偏好", "note": "战略方向 · 参会 · 协会任职"},
                {
                    "key": "external",
                    "label": "链外关键配套",
                    "note": "非上市 / 未入链的链内参与者",
                },
            ],
            "coverage": rel_coverage,
            "stats": {
                "total": relations_stats["total"],
                "supply": relations_stats["supply"],
                "equity": relations_stats["equity"],
                "companies": relations_stats["companies"],
                "profiles": relations_stats["profiles"],
                "external": external_stats["total"],
                "externalUpstream": external_stats["upstream"],
                "externalDownstream": external_stats["downstream"],
            },
        },
    ]

    payload = {
        "schemaVersion": 6,
        "generatedAt": "2026-09-14",
        "industry": INDUSTRY_LABEL,
        "industryCode": INDUSTRY_CODE,
        "source": {
            "chain": os.path.basename(SRC_CHAIN),
            "heads": os.path.basename(SRC_HEADS),
            "relations": os.path.basename(SRC_REL),
            "nodes": "Sheet1（位置 / 一级环节 / 二级环节 / 顺序）",
            "companies": "Sheet3（企业明细）",
        },
        "stats": stats,
        "positions": positions,
        "levels": levels,
        "companies": companies,
        "relations": relations,
        # 链外被点名主体：237 是上市企业样本，这一层补上非上市 / 未入链的链内参与者
        "externalMentions": external_mentions,
        "externalVerified": EXTERNAL_VERIFIED,
        # 非 237 的本链客户（来源 CNESA / CIAPS），带节点编码与属地；用于地图「全部企业」范围
        "otherCustomers": other_customers,
        "customerStats": customer_stats,
        "otherNote": (
            "这 1740 家来自 CNESA / CIAPS 的储能企业清单，按「切法B节点」编码归入同一套 24 个"
            "二级环节，属**非上市**客户口径。图谱/周报的统计仍以上市企业（237 家）为准，"
            "它们只在「全部企业」范围与地图上出现。落点方式：有本地点位缓存的用真实坐标，"
            "其余按属地**地级市**取本地底图地名点作「区域参考」——不是门牌级精确位置。"
        ),
        # 属地分布：省 / 市 / 区县 / 街道 / 园区（来自 07_属地信息，工商登记注册地址为准）
        "territory": territory_stats,
        # 全量客户属地分布（237 上市 + 1740 非上市），用于流程页「全量客户」视图
        "customerTerritory": customer_territory,
        "territoryNote": (
            "属地以工商登记「注册地址」为准，解析省 / 地级市 / 区县 / 街道 / 园区五级。"
            "区县为空的多为「园区 / 开发区 / 道路」型注册地（源地址本身未载区县），非解析失败；"
            "华润微（开曼）、华虹公司（中国香港）为离岸注册主体，境内另有运营主体，单列不入国内分布。"
        ),
        "externalNote": (
            "本链 237 家是「上市企业」样本，但产业链不只含上市公司。本层收录被这 237 家在"
            "「前五大客户 / 供应商」里具名点名、却不在 237 之内的主体 —— 既包括非上市企业"
            "（例：天齐锂业的供应商「河北吉诚」为非上市电池级碳酸锂 / 氢氧化锂厂商，"
            "属本链「资源与材料」环节），也包括已上市但未纳入本链样本的公司。"
            "关系边只画两端都在 237 内的情况，链外主体单列于此，供后续补充核实。"
        ),
        "steps": steps,
        "distributions": {
            "grades": [{"k": k, "v": v} for k, v in grade_count.most_common()],
            "fillMarks": [{"k": k, "v": v} for k, v in fill_count.most_common()],
            "boards": [{"k": k, "v": v} for k, v in board_count.most_common()],
            "swSectors": [{"k": k, "v": v} for k, v in sw_count.most_common(12)],
        },
        "warnings": {
            "emptyNodes": empty_nodes,
            "nodesWithoutHeads": nodes_without_heads,
            "unmatchedCompanyRows": [{"key": k, "count": v} for k, v in unmatched_rows.items()],
            "headNamesNotInCompanyTable": head_names_not_found,
            "headAliasApplied": alias_hits,
            "tierConflicts": conflicts,
            "relationAliasApplied": rel_alias_hits,
            "companiesWithoutRelationProfile": rel_missing,
            "relationRowsNotInCompanyTable": rel_unmatched,
            "territoryMissing": sorted(set(company_short_names) - set(territory_by_short)),
            "territoryRowsNotInCompanyTable": sorted(
                set(territory_by_short) - set(company_short_names)
            ),
            "officialNodeCodeMismatch": node_mismatch,
            "otherCustomersWithoutNode": other_no_node,
            "otherCustomersUnplaced": other_unplaced,
        },
    }

    os.makedirs(os.path.dirname(OUT_CHAIN), exist_ok=True)
    with open(OUT_CHAIN, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, separators=(",", ":"))

    # ================= 一级链数据分片 =================
    # 契约见 test/newIndustryData.test.js：companies[].classifications 非空、
    # 上市企业必须有 stockCode 且代码唯一、relations 两端必须是上市企业。
    sec_index = {lv["label"]: f"{i + 1}" for i, lv in enumerate(levels)}
    ter_index = {}
    for lv in levels:
        for ch in lv["children"]:
            ter_index[(lv["label"], ch["label"])] = f"{sec_index[lv['label']]}.{ch['order']}"

    shard_companies = []
    for i, c in enumerate(companies, start=1):
        # 落点解析结果算一次复用（真实缓存优先，其次按属地取本地底图地名点 → 区域参考）
        c_loc, c_loc_status, c_loc_precision, c_loc_source, _ = resolve_location(
            c["name"],
            (c.get("territory") or {}).get("district", ""),
            (c.get("territory") or {}).get("city", ""),
            (c.get("territory") or {}).get("province", ""),
        )
        # 该企业出现的每个 (一级, 二级) 节点各生成一条分类记录
        classifications = []
        for node in c["nodes"]:
            l1, _, l2 = node.partition(" / ")
            classifications.append(
                {
                    "chain": l1,
                    "secondaryCode": sec_index.get(l1, "0"),
                    "sector": l2,
                    "tertiaryCode": ter_index.get((l1, l2), ""),
                    "status": "mapped",
                }
            )

        primary = c["level1"][0] if c["level1"] else ""
        market = "港股" if c["board"] == "港交所" else "A股"
        shard_companies.append(
            {
                "id": f"es-{i:04d}",
                "name": c["fullName"] or c["name"],
                "normalizedName": c["fullName"] or c["name"],
                "aliases": [],
                "companyType": "ordinary",
                "isKey": False,
                "primaryIndustry": INDUSTRY_LABEL,
                "industryCode": INDUSTRY_CODE,
                "rawIndustry": primary,
                "sourceFlag": "Y",
                "chain": primary,
                "secondaryIndustry": primary,
                "secondaryCode": sec_index.get(primary, "0"),
                "sector": c["level2"][0] if c["level2"] else "",
                "tertiarySector": c["level2"][0] if c["level2"] else "",
                "tertiaryCode": ter_index.get((primary, c["level2"][0] if c["level2"] else ""), ""),
                "classifications": classifications,
                "secondaryCodes": sorted({cl["secondaryCode"] for cl in classifications}),
                "classificationStatus": "mapped",
                "classificationMethod": "source-explicit",
                "classificationConfidence": 0.95,
                "classificationEvidence": [
                    {"source": "source-workbook", "terms": [f"产业链环节：{primary}"]}
                ],
                # 属地以 07_属地信息 为准（工商登记注册地址），源表地址作为兜底。
                # 实测 237 家中仅 2 家与源表不同，且均为属地表的更正（如 广大特材
                #「凤凰镇安庆村」→「江苏省苏州市张家港市凤凰镇安庆村」）。
                "address": (c.get("territory") or {}).get("address") or c["address"],
                "addressSource": (
                    "relation-archive-territory"
                    if (c.get("territory") or {}).get("address")
                    else ("source-workbook" if c["address"] else "missing")
                ),
                "addressType": "注册地址"
                if ((c.get("territory") or {}).get("address") or c["address"])
                else "",
                "region": {
                    "province": (c.get("territory") or {}).get("province", ""),
                    "city": city_text(
                        (c.get("territory") or {}).get("province", ""),
                        (c.get("territory") or {}).get("city", ""),
                    ),
                    "district": (c.get("territory") or {}).get("district", ""),
                    "street": (c.get("territory") or {}).get("street", ""),
                },
                "regionStatus": (
                    "已从属地表解析"
                    if (c.get("territory") or {}).get("province")
                    else "待解析"
                ),
                "location": c_loc,
                "locationStatus": c_loc_status,
                "locationSource": c_loc_source,
                "locationEvidence": [],
                "coreTechnology": "",
                "coreTechnologySource": "not-researched",
                "procurement": "",
                "products": c["products"],
                "productsSource": "source-workbook",
                "suppliers": {"level1": [], "level2": []},
                "distributors": {"level1": [], "level2": []},
                "equityRaw": "",
                "controller": "",
                "beneficiary": "",
                # ⚠️ evidence 必须是**字符串数组**：CompanyDetail 直接渲染 {item}，
                # 放对象会抛 "Objects are not valid as a React child"（本次已踩）。
                "evidence": [c["evidence"]] if c["evidence"] else [],
                "tags": ["上市企业"],
                "tagStatus": "已生成",
                "tagSource": "listed-market-match",
                "bank": {
                    "isOurCustomer": None,
                    "isCreditCustomer": None,
                    "otherBankCustomer": None,
                    "sourceBatch": "not-researched",
                },
                "research": {
                    "status": "储能产业入链 + 头部标签",
                    "source": "storage-chain-workbook",
                    "confidence": 0.9,
                    "lastReviewedAt": "2026-09-14",
                    "additionalFields": {
                        "stockCode": c["code"],
                        "stockSymbol": "",
                        "stockShort": c["name"],
                        "market": market,
                        "storageGrade": c["grade"],
                        "fillMark": c["fillMark"],
                        "headTier": c["tierLabel"],
                        "headSegments": c["tierSegments"],
                        "relationCount": c["relationCount"],
                        # 属地（来自 07_属地信息），详情栏直接展示
                        "territoryProvince": (c.get("territory") or {}).get("province", ""),
                        "territoryCity": (c.get("territory") or {}).get("city", ""),
                        "territoryDistrict": (c.get("territory") or {}).get("district", ""),
                        "territoryStreet": (c.get("territory") or {}).get("street", ""),
                        "territoryPark": (c.get("territory") or {}).get("park", ""),
                        "territoryAddress": (c.get("territory") or {}).get("address", ""),
                        "territoryNote": (c.get("territory") or {}).get("note", ""),
                    },
                },
                "profileId": "",
            }
        )

    # ---------- 非 237 客户写进分片（非上市口径，供地图「全部企业」范围使用） ----------
    # ⚠️ 这些企业 **不带** `上市企业` tag —— 地图/图谱/周报的统计口径仍以上市企业为准，
    #    它们只在「全部企业」范围里出现，不改变 237 家的任何指标。
    for j, item in enumerate(other_customers, start=1):
        classifications = []
        for code in item["codes"]:
            pair = code_to_node.get(code)
            if not pair:
                continue
            l1, l2 = pair
            classifications.append(
                {
                    "chain": l1,
                    "secondaryCode": sec_index.get(l1, "0"),
                    "sector": l2,
                    "tertiaryCode": ter_index.get((l1, l2), ""),
                    "status": "mapped",
                }
            )
        if not classifications:
            # 理论上不该出现（1740 家全部有节点编码）；兜底成「待研判」，不让它挂在空分类上
            classifications.append(
                {
                    "chain": "",
                    "secondaryCode": "0",
                    "sector": "",
                    "tertiaryCode": "",
                    "status": "unmapped",
                }
            )
        primary_pair = code_to_node.get(item["primaryCode"]) or ("", "")
        primary_l1, primary_l2 = primary_pair
        shard_companies.append(
            {
                "id": f"es-o{j:04d}",
                "name": item["name"],
                "normalizedName": item["name"],
                "aliases": [],
                "companyType": "ordinary",
                "isKey": False,
                "primaryIndustry": INDUSTRY_LABEL,
                "industryCode": INDUSTRY_CODE,
                "rawIndustry": primary_l1,
                "sourceFlag": "N",
                "chain": primary_l1,
                "secondaryIndustry": primary_l1,
                "secondaryCode": sec_index.get(primary_l1, "0"),
                "sector": primary_l2,
                "tertiarySector": primary_l2,
                "tertiaryCode": ter_index.get((primary_l1, primary_l2), ""),
                "classifications": classifications,
                "secondaryCodes": sorted({cl["secondaryCode"] for cl in classifications}),
                "classificationStatus": "mapped",
                "classificationMethod": "source-explicit",
                "classificationConfidence": 0.9,
                "classificationEvidence": [
                    {"source": "all-customer-workbook", "terms": [f"切法B节点：{item['primaryCode']}"]}
                ],
                "address": item["address"],
                "addressSource": "all-customer-workbook",
                "addressType": "注册地址" if item["address"] else "",
                "region": {
                    "province": item["province"],
                    "city": city_text(item["province"], item["city"]),
                    "district": item["district"] if item["district"] != "—" else "",
                    "street": "",
                },
                "regionStatus": "已从客户清单解析" if item["province"] else "待解析",
                "location": item["location"],
                "locationStatus": item["locationStatus"],
                "locationSource": item["locationSource"],
                "locationEvidence": [],
                "coreTechnology": "",
                "coreTechnologySource": "not-researched",
                "procurement": "",
                "products": item["mainBusiness"],
                "productsSource": "all-customer-workbook" if item["mainBusiness"] else "not-researched",
                "suppliers": {"level1": [], "level2": []},
                "distributors": {"level1": [], "level2": []},
                "equityRaw": "",
                "controller": "",
                "beneficiary": "",
                "evidence": [],
                # 刻意不带 `上市企业`，避免影响任何以上市企业为口径的统计
                "tags": [],
                "tagStatus": "未标注",
                "tagSource": "all-customer-list",
                "bank": {
                    "isOurCustomer": None,
                    "isCreditCustomer": None,
                    "otherBankCustomer": None,
                    "sourceBatch": "not-researched",
                },
                "research": {
                    "status": "储能产业非上市客户",
                    "source": "all-customer-workbook",
                    "confidence": 0.85,
                    "lastReviewedAt": "2026-09-15",
                    "additionalFields": {
                        "stockCode": "",
                        "stockSymbol": "",
                        "stockShort": item["name"],
                        "market": "",
                        "creditCode": item["creditCode"],
                        "memberSource": item["source"],
                        "verifyStatus": item["verifyStatus"],
                        "primaryNodeCode": item["primaryCode"],
                        "nodeCodes": item["codes"],
                        "nodeDetail": item["nodeDetail"],
                        "territoryProvince": item["province"],
                        "territoryCity": item["city"],
                        "territoryDistrict": item["district"] if item["district"] != "—" else "",
                        "territoryAddress": item["address"],
                        "locationMatchedName": item["locationMatchedName"],
                    },
                },
                "profileId": "",
            }
        )

    shard = {
        "generatedAt": "2026-09-14",
        "schemaVersion": 2,
        "industry": INDUSTRY_LABEL,
        "industryCode": INDUSTRY_CODE,
        "counts": {
            "totalCompanies": len(shard_companies),
            "mapped": 0,
            "verifiedLocation": sum(
                1 for c in shard_companies if c["locationStatus"] in ("已复核", "已核实")
            ),
            "candidateLocation": sum(
                1 for c in shard_companies if c["locationStatus"] == "候选"
            ),
            "regionReferenceLocation": sum(
                1 for c in shard_companies if c["locationStatus"] == "区域参考"
            ),
            "pendingLocation": sum(1 for c in shard_companies if not c["location"]),
            "missingAddress": sum(1 for c in shard_companies if not c["address"]),
            "listed": len(companies),
            "nonListed": len(other_customers),
            "neeqListed": 0,
            "relationCount": len(relations),
            "classification": {"secondary": {lv["label"]: lv["companyCount"] for lv in levels}},
            "headTiers": {
                "龙头": stats["leader"],
                "骨干": stats["core"],
                "长尾": stats["tail"],
            },
            "relationTypes": {
                "供应": relations_stats["supply"],
                "股权": relations_stats["equity"],
            },
        },
        "companies": shard_companies,
        "relations": relations,
    }
    with open(OUT_SHARD, "w", encoding="utf-8") as fh:
        json.dump(shard, fh, ensure_ascii=False, separators=(",", ":"))

    # 契约定检：CompanyDetail 会直接把 evidence 的每一项当 React child 渲染，
    # 因此必须是标量；这里挡住"改成对象"这类回归。
    for item in shard_companies:
        for value in item["evidence"]:
            if not isinstance(value, str):
                raise TypeError(f"evidence 必须是字符串数组，{item['name']} 出现 {type(value).__name__}")
        for value in item["tags"]:
            if not isinstance(value, str):
                raise TypeError(f"tags 必须是字符串数组，{item['name']} 出现 {type(value).__name__}")

    # ---------- 输出 ----------
    print(f"✓ {OUT_CHAIN}  ({os.path.getsize(OUT_CHAIN) / 1024 / 1024:.2f} MB)")
    print(f"✓ {OUT_SHARD}  ({os.path.getsize(OUT_SHARD) / 1024 / 1024:.2f} MB)")
    print(
        f"  一级环节 {stats['level1']} / 二级环节 {stats['level2']} / "
        f"企业 {stats['companies']}（龙头 {stats['leader']} · 骨干 {stats['core']} · 长尾 {stats['tail']}）"
    )
    print()
    for lv in levels:
        print(
            f"  {lv['code']} {lv['label']} [{lv['position']}] {lv['childCount']}个二级 · "
            f"{lv['companyCount']}家 · 龙{lv['leaderCount']}/骨{lv['coreCount']}/尾{lv['tailCount']}"
        )
        for ch in lv["children"]:
            flag = ""
            if ch["companyCount"] == 0:
                flag = " ⚠️0家"
            elif not ch["heads"]:
                flag = " ⚠️无头部名单"
            segs = "/".join(h["segment"] for h in ch["heads"])
            print(
                f"       {ch['label']}: {ch['companyCount']}家 · "
                f"龙{ch['leaderCount']}/骨{ch['coreCount']}/尾{ch['tailCount']}"
                f"{' · ' + segs if segs else ''}{flag}"
            )
    print()
    print(
        f"  关系：{relations_stats['total']} 条（供销 {relations_stats['supply']} · "
        f"股权 {relations_stats['equity']}）覆盖 {relations_stats['companies']} 家企业；"
        f"关系档案覆盖 {relations_stats['profiles']}/{len(company_short_names)} 家"
    )
    print(
        f"  全量客户：上市 {customer_stats['listed']} + 非上市 {customer_stats['other']} = "
        f"{customer_stats['total']} 家；落点 真实坐标 {customer_stats['otherRealLocation']} · "
        f"区域参考 {customer_stats['otherRegionReference']} · 未落图 {customer_stats['otherUnplaced']}"
    )
    print(
        "     非上市客户来源："
        + " / ".join(f"{x['k']} {x['v']}" for x in customer_stats["otherBySource"])
        + "；核验状态："
        + " / ".join(f"{x['k']} {x['v']}" for x in customer_stats["otherByVerify"])
    )
    if node_mismatch:
        print(f"  ⚠️ 节点关系表与主表不一致：{node_mismatch}")
    print(
        f"  链外被点名主体：{external_stats['total']} 个"
        f"（偏上游 {external_stats['upstream']} · 偏下游 {external_stats['downstream']}，"
        f"门槛 >= {EXTERNAL_MIN_COUNT} 次）"
    )
    for item in external_mentions[:12]:
        who = "、".join(m["company"] for m in item["mentionedBy"][:4])
        print(f"     {item['name']}: {item['count']} 次（上{item['asUpstream']}/下{item['asDownstream']}）← {who}")
    print(
        f"  属地：{territory_stats['total']} 家覆盖 {territory_stats['provinces']} 个省级行政区 / "
        f"{territory_stats['cities']} 个地级市；区县 {territory_stats['withDistrict']} 家、"
        f"街道 {territory_stats['withStreet']} 家、园区 {territory_stats['withPark']} 家；"
        f"离岸注册 {territory_stats['offshore']} 家"
    )
    print(
        "     省级 Top："
        + " / ".join(f"{p['k']} {p['v']}" for p in territory_stats["topProvinces"][:8])
    )
    print(
        "     城市 Top："
        + " / ".join(f"{p['k']} {p['v']}" for p in territory_stats["topCities"][:8])
    )
    for cov in rel_coverage:
        print(f"     {cov['key']}: 完整 {cov['full']} / 部分 {cov['partial']} / 缺失 {cov['missing']}")
    print()
    if alias_hits:
        print(f"  别名应用（头部）：{alias_hits}")
    if rel_alias_hits:
        print(f"  别名应用（关系档案）：{rel_alias_hits}")
    if head_names_not_found:
        print(f"  ⚠️ 头部名单中不在企业表的名字：{head_names_not_found}")
    if rel_missing:
        print(f"  ⚠️ 无关系档案记录的企业：{rel_missing}")
    if rel_unmatched:
        print(f"  ⚠️ 关系档案中不在本链企业表的名字：{rel_unmatched}")
    if empty_nodes:
        print(f"  ⚠️ 无企业的二级环节：{empty_nodes}")
    if nodes_without_heads:
        print(f"  ⚠️ 无头部名单的二级环节（企业全记长尾）：{nodes_without_heads}")
    if conflicts:
        print(f"  ⚠️ 档位冲突：{conflicts}")


if __name__ == "__main__":
    main()
